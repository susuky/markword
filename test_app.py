import os
import time
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app import _annotate_html_lines, _cleanup_old_exports, _get_export_filename, _render_md_to_html_for_export, analyze_text, export_pdf, export_word, render_preview, EXPORT_DIR
from themes import THEMES


class TestAnalyzeText(unittest.TestCase):
    '''
    Test cases for text analysis metrics calculation in analyze_text function.
    '''

    def test_empty_string(self):
        '''
        Test analyzing an empty string.
        '''
        res = analyze_text('')
        self.assertEqual(res, (0, 0, 0, 0, 0, 0, 0))

    def test_chinese_text(self):
        '''
        Test analyzing pure Chinese text with punctuation.
        '''
        text = '你好世界！這是測試。'
        total, no_space, cjk, cjk_punct, eng, digits, lines = analyze_text(text)
        self.assertEqual(cjk, 8)
        self.assertEqual(cjk_punct, 2)
        self.assertEqual(eng, 0)
        self.assertEqual(lines, 1)

    def test_mixed_text(self):
        '''
        Test analyzing mixed Chinese, English, digits, and punctuation.
        '''
        text = 'Hello World 2026! 你好 123.\n第二行測試。'
        total, no_space, cjk, cjk_punct, eng, digits, lines = analyze_text(text)
        self.assertEqual(cjk, 7)
        self.assertEqual(cjk_punct, 1)
        self.assertEqual(eng, 2)
        self.assertEqual(digits, 7)
        self.assertEqual(lines, 2)


class TestExportFilename(unittest.TestCase):
    '''
    Test cases for export filename determination (_get_export_filename).
    '''

    def test_h1_header(self):
        '''
        Test that filename uses H1 header if present.
        '''
        md_text = '# 我的專案報告\n\n這是內容'
        filename = _get_export_filename(md_text, 'pdf')
        self.assertEqual(filename, '我的專案報告.pdf')

    def test_h1_with_formatting(self):
        '''
        Test H1 header with markdown formatting characters cleaned.
        '''
        md_text = '# **重點** *說明*\n\n這是內容'
        filename = _get_export_filename(md_text, 'docx')
        self.assertEqual(filename, '重點 說明.docx')

    def test_no_h1_header(self):
        '''
        Test fallback to export_{hash} when no H1 header is present.
        '''
        md_text = '## H2 標題\n沒有 H1 標題'
        filename = _get_export_filename(md_text, 'pdf')
        self.assertTrue(filename.startswith('export_'))
        self.assertTrue(filename.endswith('.pdf'))

    def test_h1_filename_is_bounded(self):
        filename = _get_export_filename(f'# {"長" * 300}', 'pdf')
        self.assertLessEqual(len(filename.encode('utf-8')), 240)
        self.assertTrue(filename.endswith('.pdf'))

    def test_long_unicode_titles_are_written_in_both_formats(self):
        for title in ['長' * 120, 'Report台灣🌟' * 60]:
            with tempfile.TemporaryDirectory() as directory, patch('app.EXPORT_DIR', directory):
                for exporter in [export_pdf, export_word]:
                    path = exporter(f'# {title}\n\n保留全文')
                    self.assertTrue(os.path.isfile(path))
                    self.assertLessEqual(len(os.path.basename(path).encode('utf-8')), 240)

    def test_cleanup_old_exports(self):
        '''
        Test that _cleanup_old_exports removes old temporary export files.
        '''
        dummy_dir = os.path.join(EXPORT_DIR, 'test_old_dir')
        os.makedirs(dummy_dir, exist_ok=True)
        # Set mtime to 1 hour ago
        old_time = time.time() - 3600
        os.utime(dummy_dir, (old_time, old_time))

        _cleanup_old_exports(max_age_seconds=600)
        self.assertFalse(os.path.exists(dummy_dir))


class TestRenderPreview(unittest.TestCase):
    '''
    Test cases for markdown preview rendering and themes.
    '''

    def test_empty_input(self):
        '''
        Test rendering empty input shows placeholder message.
        '''
        result = render_preview('')
        self.assertIn('請在左側輸入', result)

    def test_heading_render(self):
        '''
        Test that markdown heading is rendered to HTML h1 tag inside iframe.
        '''
        result = render_preview('# 測試標題', 'Light')
        self.assertIn('iframe', result)
        self.assertIn('&lt;h1', result)
        self.assertIn('測試標題', result)

    def test_themes_exist(self):
        '''
        Test that all defined themes can render without error.
        '''
        for theme_name in ['Light', 'Paper', 'Sage', 'Dark', 'Ocean', 'Nord', 'Dracula', 'Midnight']:
            self.assertIn(theme_name, THEMES)
            result = render_preview('# 測試標題', theme_name)
            self.assertIn('iframe', result)

    def test_codehilite_syntax_highlighting(self):
        '''
        Test that code blocks produce Pygments syntax highlight spans.
        '''
        md_code = '```python\ndef foo():\n    return 42\n```'
        result = render_preview(md_code, 'Dark')
        self.assertIn('codehilite', result)

    def test_data_line_annotation(self):
        '''
        Test that _annotate_html_lines adds data-line attributes to HTML tags.
        '''
        md_text = '# 標題1\n\n段落內容'
        body_html = '<h1>標題1</h1>\n<p>段落內容</p>'
        annotated = _annotate_html_lines(md_text, body_html)
        self.assertIn('data-line="1"', annotated)

    def test_export_pdf_and_word_themes(self):
        '''
        Test PDF and Word exports with themes and custom filename.
        '''
        md_text = '# 標題\n\n```python\nprint("hello")\n```'
        for theme_name in ['Light', 'Dark']:
            pdf_path = export_pdf(md_text, theme_name)
            word_path = export_word(md_text, theme_name)
            self.assertIsNotNone(pdf_path)
            self.assertIsNotNone(word_path)
            self.assertTrue(pdf_path.endswith('標題.pdf'))
            self.assertTrue(word_path.endswith('標題.docx'))

    def test_export_layout_css_is_applied(self):
        html = _render_md_to_html_for_export('# 文件標題\n\n內容', THEMES['Paper'], 'Editorial')
        self.assertIn('Noto Serif CJK TC', html)
        self.assertIn('@page', html)
        self.assertIn('#fbf7ef', html)

    def test_table_export_preserves_line_breaks_pipes_and_alignment(self):
        markdown = '| 項目 | 備註 |\n| :--- | ---: |\n| A\\|B | 第一行<br>第二行 |'
        with tempfile.TemporaryDirectory() as directory, patch('app.EXPORT_DIR', directory):
            document = Document(export_word(markdown))
        table = document.tables[0]
        self.assertEqual(table.cell(1, 0).text, 'A|B')
        self.assertEqual(table.cell(1, 1).text, '第一行\n第二行')
        self.assertEqual(table.cell(1, 0).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(table.cell(1, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        html = _render_md_to_html_for_export(markdown, THEMES['Light'])
        self.assertIn('第一行<br>第二行', html)
        self.assertIn('A|B', html)

    def test_word_nested_lists_keep_text_order_indentation_and_numbering(self):
        markdown = '3. Parent\n    - Child **bold** *italic*\n        7. Grandchild\n        8. Next grandchild\n    - Next child\n4. Next parent\n\nParagraph\n\n1. Restart'
        with tempfile.TemporaryDirectory() as directory, patch('app.EXPORT_DIR', directory):
            document = Document(export_word(markdown))
        paragraphs = document.paragraphs
        self.assertEqual([p.text.strip() for p in paragraphs], [
            'Parent', 'Child bold italic', 'Grandchild', 'Next grandchild', 'Next child', 'Next parent', 'Paragraph', 'Restart',
        ])
        self.assertEqual([p.paragraph_format.left_indent.inches for p in paragraphs[:6]], [.25, .5, .75, .75, .5, .25])
        self.assertEqual([p._p.pPr.numPr.ilvl.val for p in paragraphs[:6]], [0, 1, 2, 2, 1, 0])
        self.assertTrue(any(run.bold and run.text == 'bold' for run in paragraphs[1].runs))
        numbering = document.part.numbering_part.element
        def numbering_values(paragraph):
            num_id = paragraph._p.pPr.numPr.numId.val
            num = numbering.xpath(f'w:num[@w:numId="{num_id}"]')[0]
            abstract_id = num.find(qn('w:abstractNumId')).get(qn('w:val'))
            level = numbering.xpath(f'w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl')[0]
            return num_id, level.find(qn('w:start')).get(qn('w:val')), level.find(qn('w:numFmt')).get(qn('w:val'))
        parent = numbering_values(paragraphs[0])
        self.assertEqual(parent[1:], ('3', 'decimal'))
        self.assertEqual(numbering_values(paragraphs[5]), parent)
        self.assertEqual(numbering_values(paragraphs[1])[2], 'bullet')
        self.assertEqual(numbering_values(paragraphs[2])[1:], ('7', 'decimal'))
        self.assertNotEqual(numbering_values(paragraphs[7])[0], parent[0])
        self.assertEqual(numbering_values(paragraphs[7])[1:], ('1', 'decimal'))


if __name__ == '__main__':
    unittest.main()
