import os
from pathlib import Path

import httpx
import pytest

import backend.main as main


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
async def client():
    transport = httpx.ASGITransport(app=main.app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as api_client:
        yield api_client


@pytest.mark.anyio
async def test_health(client):
    response = await client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


@pytest.mark.anyio
async def test_analyze_mixed_text(client):
    response = await client.post("/api/analyze", json={"text": "Hello 台灣 123\n第二行"})
    assert response.status_code == 200
    assert response.json() == {
        "total_chars": 16,
        "chars_no_spaces": 13,
        "cjk_count": 5,
        "cjk_punct_count": 0,
        "english_words": 1,
        "digit_count": 3,
        "line_count": 2,
    }


@pytest.mark.anyio
async def test_themes(client):
    response = await client.get("/api/themes")
    assert response.status_code == 200
    payload = response.json()
    assert payload["default_theme"] == "Light"
    assert {theme["name"] for theme in payload["themes"]} == {
        "Light",
        "Dark",
        "Nord",
        "Dracula",
        "Paper",
        "Sage",
        "Ocean",
        "Midnight",
    }


@pytest.mark.anyio
async def test_export_rejects_blank_markdown(client):
    response = await client.post(
        "/api/export/pdf",
        json={"markdown": "   ", "theme": "Light"},
    )
    assert response.status_code == 422


@pytest.mark.anyio
async def test_export_rejects_unknown_theme(client):
    response = await client.post(
        "/api/export/docx",
        json={"markdown": "# Test", "theme": "Missing"},
    )
    assert response.status_code == 422


@pytest.mark.anyio
async def test_export_rejects_unknown_style(client):
    response = await client.post(
        "/api/export/pdf",
        json={"markdown": "# Test", "theme": "Light", "style": "Missing"},
    )
    assert response.status_code == 422


@pytest.mark.anyio
@pytest.mark.parametrize('setting', [
    {'page_size': 'A4; margin:0'}, {'margin_mm': 0}, {'margin_mm': 41},
    {'font_size': 11}, {'font_size': 25},
])
@pytest.mark.parametrize('export_format', ['pdf', 'docx'])
async def test_export_rejects_invalid_page_settings(client, setting, export_format):
    response = await client.post(f'/api/export/{export_format}', json={'markdown': '# Page', **setting})
    assert response.status_code == 422


@pytest.mark.anyio
@pytest.mark.parametrize('export_format', ['pdf', 'docx'])
async def test_worker_applies_page_settings(client, monkeypatch, tmp_path, export_format):
    from io import BytesIO
    from docx import Document
    import re
    import zlib
    import app as export_app

    monkeypatch.setattr(export_app, 'EXPORT_DIR', str(tmp_path))
    response = await client.post(f'/api/export/{export_format}', json={
        'markdown': '# Page settings\n\nA readable paragraph with **bold** and [a link](https://example.com).',
        'page_size': 'Letter', 'margin_mm': 24, 'font_size': 20,
    })
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    if export_format == 'pdf':
        objects = response.content
        for stream in re.findall(rb'stream\r?\n(.*?)\r?\nendstream', response.content, re.S):
            try:
                objects += zlib.decompress(stream)
            except zlib.error:
                pass
        box = re.search(rb'/MediaBox\s*\[([^]]+)\]', objects)
        assert box and [float(value) for value in box[1].split()] == pytest.approx([0, 0, 612, 792])
    else:
        doc = Document(BytesIO(response.content))
        section = doc.sections[0]
        assert section.page_width.inches == pytest.approx(8.5)
        assert section.page_height.inches == pytest.approx(11)
        assert section.left_margin.mm == pytest.approx(24, abs=.02)
        assert section.top_margin.mm == pytest.approx(24, abs=.02)
        assert doc.styles['Normal'].font.size.pt == 15
        body_sizes = [int(size.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) for size in doc.paragraphs[1]._p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')]
        assert body_sizes and set(body_sizes) == {30}
    assert list(tmp_path.iterdir()) == []


@pytest.mark.anyio
async def test_pdf_export_download_and_cleanup(client, monkeypatch, tmp_path):
    export_dir = tmp_path / "pdf_request"
    export_dir.mkdir()
    export_path = export_dir / "測試文件.pdf"
    export_path.write_bytes(b"%PDF-test")

    seen = {}

    async def fake_export(payload, export_format):
        seen.update(markdown=payload.markdown, theme=payload.theme, style=payload.style)
        assert export_format == 'pdf'
        return str(export_path), str(export_dir)

    monkeypatch.setattr(main, "run_export", fake_export)
    response = await client.post(
        "/api/export/pdf",
        json={"markdown": "# 測試文件", "theme": "Light"},
    )

    assert response.status_code == 200
    assert response.content == b"%PDF-test"
    assert response.headers["content-type"] == "application/pdf"
    assert "filename*=utf-8''" in response.headers["content-disposition"].lower()
    assert seen == {"markdown": "# 測試文件", "theme": "Light", "style": "Classic"}
    assert not export_dir.exists()


@pytest.mark.anyio
async def test_pdf_export_rejects_local_resources_without_disclosing_path(client, monkeypatch, tmp_path):
    import app as export_app

    monkeypatch.setattr(export_app, "EXPORT_DIR", str(tmp_path))
    response = await client.post(
        "/api/export/pdf",
        json={"markdown": '<link rel="attachment" href="file:///tmp/private.txt">'},
    )
    assert response.status_code == 422
    assert "embedded" in response.json()["detail"]
    assert "private.txt" not in response.text
    assert list(tmp_path.iterdir()) == []


@pytest.mark.anyio
@pytest.mark.parametrize('style', ['Editorial', 'Report', 'Compact'])
async def test_word_rejects_unsupported_layouts(client, style):
    response = await client.post('/api/export/docx', json={'markdown': '# Word', 'style': style})
    assert response.status_code == 422


@pytest.mark.anyio
@pytest.mark.parametrize('export_format', ['pdf', 'docx'])
async def test_real_export_job_downloads_and_cleans_up(client, monkeypatch, tmp_path, export_format):
    from io import BytesIO
    from docx import Document
    import app as export_app

    monkeypatch.setattr(export_app, 'EXPORT_DIR', str(tmp_path))
    response = await client.post(f'/api/export/{export_format}', json={
        'markdown': '# 匯出🌟' + '長標題' * 60 + '\n\n1. Parent\n    - Child\n\n```mermaid\nflowchart LR\n A[開始] --> B[完成]\n```',
    })
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    assert "filename*=utf-8''" in response.headers['content-disposition'].lower()
    if export_format == 'pdf':
        assert response.content.startswith(b'%PDF-')
        assert b'/Subtype /Image' in response.content
    else:
        doc = Document(BytesIO(response.content))
        assert any(p.text.strip() == 'Child' for p in doc.paragraphs)
        assert len(doc.inline_shapes) == 1
    assert list(tmp_path.iterdir()) == []


@pytest.mark.anyio
@pytest.mark.parametrize('export_format', ['pdf', 'docx'])
async def test_worker_preserves_math_and_embedded_image(client, monkeypatch, tmp_path, export_format):
    import base64
    from io import BytesIO
    from docx import Document
    from PIL import Image
    import app as export_app

    monkeypatch.setattr(export_app, 'EXPORT_DIR', str(tmp_path))
    image = BytesIO()
    Image.new('RGB', (20, 20), '#237474').save(image, format='PNG')
    image_url = 'data:image/png;base64,' + base64.b64encode(image.getvalue()).decode()
    response = await client.post(f'/api/export/{export_format}', json={
        'markdown': f'# Formula and picture\n\nBefore $x^2$ after.\n\n![Photo]({image_url})\n\n- [x] Done',
    })
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    if export_format == 'pdf':
        assert response.content.startswith(b'%PDF-')
        assert response.content.count(b'/Subtype /Image') >= 2
    else:
        doc = Document(BytesIO(response.content))
        assert len(doc.inline_shapes) == 2
        assert any(p.text == 'Before  after.' for p in doc.paragraphs)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.anyio
async def test_spa_fallback(client, monkeypatch, tmp_path):
    (tmp_path / "index.html").write_text("<main>Markword</main>", encoding="utf-8")
    monkeypatch.setattr(main, "FRONTEND_DIR", Path(tmp_path))

    response = await client.get("/editor/document-1")
    assert response.status_code == 200
    assert "Markword" in response.text


@pytest.mark.anyio
async def test_unknown_api_route_stays_json_404(client):
    response = await client.get("/api/not-a-route")
    assert response.status_code == 404
    assert response.json() == {"detail": "Not Found"}
