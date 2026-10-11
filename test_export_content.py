"""Cross-format content and parser boundaries for the real export pipeline."""

import base64
from io import BytesIO
from pathlib import Path
import subprocess
import zipfile

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
import pytest
from weasyprint.urls import FatalURLFetchingError

import app
import export_markdown
import math_renderer
from themes import THEMES


def image_url(format='PNG', mime='image/png', size=(40, 20)):
    output = BytesIO()
    Image.new('RGB', size, '#286f74').save(output, format=format)
    return f'data:{mime};base64,' + base64.b64encode(output.getvalue()).decode()


def test_syntax_preserves_literal_code_escapes_and_unclosed_math(monkeypatch):
    formulas, diagrams = [], []
    def render_math(items, _color):
        formulas.extend(items)
        return [None] * len(items)
    def render_mermaid(source, _theme):
        diagrams.append(source)
        return None
    monkeypatch.setattr(export_markdown, 'render_math_pngs', render_math)
    monkeypatch.setattr(export_markdown, 'render_mermaid_png', render_mermaid)
    source = r'''Escaped \$x$; inline `$code$`; formula $a\{b\}$; price $5 and $10.

~~~~text
```mermaid
not a diagram
```
$literal$
~~~~~

> ~~~MERMAID description
> flowchart LR
> A --> B
> ~~~~~

- nested

    ````mermaid
    flowchart LR
    B --> C
    `````

    $$nested_math$$

$$
\frac{1}{2}
$$

$$unfinished

Final paragraph.

    $indented_literal$
'''
    soup = BeautifulSoup(export_markdown.render_export_body(source, THEMES['Light']), 'html.parser')
    assert formulas == [(r'a\{b\}', False), ('nested_math', True), (r'\frac{1}{2}', True)]
    assert diagrams == ['flowchart LR\nA --> B', 'flowchart LR\nB --> C']
    assert '$code$' in soup.code.text
    assert 'not a diagram' in soup.select_one('.codehilite').text
    assert '$literal$' in soup.select_one('.codehilite').text
    assert '$indented_literal$' in soup.text
    assert '$x$' in soup.text and '$5 and $10' in soup.text
    assert '$$unfinished' in soup.text and 'Final paragraph.' in soup.text
    assert 'Formula could not be rendered' in soup.text


@pytest.mark.parametrize('fence,closing', [('```', '```'), ('~~~', '~~~'), ('````', '`````'), ('~~~~', '~~~~~~'), ('```', '')])
def test_legal_mermaid_fences_are_rendered_once(fence, closing, monkeypatch):
    seen = []
    monkeypatch.setattr(export_markdown, 'render_mermaid_png', lambda code, _theme: seen.append(code))
    html = export_markdown.render_export_body(f'{fence}mermaid\nflowchart LR\nA --> B\n{closing}', THEMES['Light'])
    assert seen == ['flowchart LR\nA --> B']
    assert 'Mermaid diagram could not be rendered' in html


def test_footnote_fences_keep_diagrams_and_literal_examples_in_both_formats(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = '''Reference[^note].

[^note]: Before diagram.

    ~~~mermaid
    flowchart LR
    A --> B
    ~~~~

    After diagram.

    ````text
    ```mermaid
    literal example
    ```
    ````
'''
    body = BeautifulSoup(export_markdown.render_export_body(source, THEMES['Light']), 'html.parser')
    footnote = body.select_one('.footnote li')
    assert len(footnote.select('.mermaid-img img')) == 1
    assert '```mermaid\nliteral example\n```' in footnote.select_one('.codehilite').text
    assert not footnote.select('p > div, p > pre, p > p')
    doc = Document(app.export_word(source))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1
    assert text.index('Before diagram.') < text.index('After diagram.') < text.index('literal example')
    pdf = Path(app.export_pdf(source))
    pdf_text = subprocess.check_output(['pdftotext', str(pdf), '-'], text=True)
    assert 'literal example' in pdf_text and 'After diagram.' in pdf_text
    assert b'/Subtype /Image' in pdf.read_bytes()


def test_cross_format_document_preserves_formulas_footnotes_tasks_and_images(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = r'''# Export conformance

Before $x^2+y^2$ after. Footnote second[^b] then first[^a].

$$
\frac{a}{b} = \sqrt{2}
$$

- [x] Done
- [ ] Todo
- \[x] Literal checkbox

![Inline image](IMAGE) after picture.

[Safe link](https://example.invalid/guide)

~~~mermaid
flowchart LR
 A[Start] --> B[Finish]
~~~~

[^a]: Alpha note.
[^b]: Beta note with **bold** text.
'''.replace('IMAGE', image_url('WEBP', 'image/webp'))
    pdf = Path(app.export_pdf(source))
    docx = Path(app.export_word(source))
    pdf_text = subprocess.check_output(['pdftotext', str(pdf), '-'], text=True)
    assert all(text in pdf_text for text in ['Before', 'after.', 'Beta note', 'Alpha note', '☑ Done', '☐ Todo', '[x] Literal'])
    assert pdf_text.index('Beta note') < pdf_text.index('Alpha note')
    assert pdf.read_bytes().count(b'/Subtype /Image') >= 4
    document = Document(docx)
    text = '\n'.join(p.text for p in document.paragraphs)
    assert all(value in text for value in ['Before  after.', '☑ Done', '☐ Todo', '[x] Literal', 'Beta note', 'Alpha note', 'after picture.'])
    assert len(document.inline_shapes) == 4
    descriptions = [shape._inline.docPr.get('descr') for shape in document.inline_shapes]
    assert '$x^2+y^2$' in descriptions
    inline_math = document.inline_shapes[0]
    assert .2 < inline_math.width.inches < 1
    assert inline_math.height.inches < .5
    assert 'Inline image' in descriptions
    assert any(run.font.superscript and run.text == '1' for p in document.paragraphs for run in p.runs)
    assert text.index('Beta note') < text.index('Alpha note')
    with zipfile.ZipFile(docx) as archive:
        xml = archive.read('word/document.xml')
        relationships = archive.read('word/_rels/document.xml.rels')
    assert b'<w:position' in xml  # formula baseline retained, not full-width pictures
    assert b'https://example.invalid/guide' in relationships
    assert b'TargetMode="External"' in relationships


@pytest.mark.parametrize('format,mime', [('PNG', 'image/png'), ('JPEG', 'image/jpeg'), ('GIF', 'image/gif'), ('WEBP', 'image/webp'), ('BMP', 'image/bmp')])
def test_supported_images_are_embedded_in_word_paragraphs_tables_and_block_nodes(format, mime, monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    image = image_url(format, mime)
    source = f'Before ![Photo]({image}) after.\n\n| Cell |\n| --- |\n| ![Table]({image}) |\n\n<img alt="Block" src="{image}">'
    doc = Document(app.export_word(source))
    assert len(doc.inline_shapes) == 3
    assert doc.paragraphs[0].text == 'Before  after.'
    assert len(doc.tables) == 1
    assert [shape._inline.docPr.get('descr') for shape in doc.inline_shapes] == ['Photo', 'Table', 'Block']


@pytest.mark.parametrize('source', [
    'https://example.invalid/image.png',
    'file:///tmp/private.png',
    'assets/missing.png',
    'data:image/png;base64,not-base64',
    'data:image/png;base64,' + base64.b64encode(b'<svg xmlns="http://www.w3.org/2000/svg"/>').decode(),
    'data:image/svg+xml;base64,' + base64.b64encode(b'<svg xmlns="http://www.w3.org/2000/svg"/>').decode(),
    image_url(size=(8193, 1)),
])
def test_invalid_or_external_images_fail_both_formats_without_partial_files(source, monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    for exporter in (app.export_pdf, app.export_word):
        with pytest.raises(FatalURLFetchingError):
            exporter(f'![image]({source})')
    assert list(tmp_path.iterdir()) == []


def test_word_only_creates_safe_link_relationships(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = '\n\n'.join(f'<a href="{url}">Label {index}</a>' for index, url in enumerate([
        'https://example.invalid/', 'mailto:writer@example.invalid', 'file:///private.txt',
        'javascript:alert(1)', 'data:text/html,hello', './assets/attachment.pdf',
    ]))
    path = app.export_word(source)
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert './assets/attachment.pdf' in text and 'file:///private.txt' in text
    links = [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
    assert links == ['https://example.invalid/', 'mailto:writer@example.invalid']


def test_invalid_and_over_budget_math_preserve_source_and_missing_runtime_fails(monkeypatch, tmp_path):
    monkeypatch.setattr(math_renderer, 'MAX_FORMULAS', 1)
    html = export_markdown.render_export_body(r'$\unknowncommand$ and $x^2$', THEMES['Light'])
    assert '$\\unknowncommand$' in html and '$x^2$' in html
    assert html.count('Formula could not be rendered') == 2
    monkeypatch.setattr(math_renderer, 'KATEX_DIST', tmp_path)
    with pytest.raises(RuntimeError, match='Local KaTeX bundle missing'):
        export_markdown.render_export_body('$x$', THEMES['Light'])
    # Ordinary exports remain independent of the browser math runtime.
    assert '<p>Plain text</p>' == export_markdown.render_export_body('Plain text', THEMES['Light'])


def test_closed_nested_toggles_export_all_children_and_preserve_literal_tags(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = '''<details>
<summary>外層 **標題**</summary>

Visible parent text.

<details>
<summary>內層標題</summary>

Hidden child content. [Linked][ref]

```text
</details>
<summary>literal</summary>
```

| 項目 | 內容 |
| --- | --- |
| 資料 | 完整保留 |

</details>

</details>

[ref]: https://example.invalid/guide
'''
    body = BeautifulSoup(export_markdown.render_export_body(source, THEMES['Light']), 'html.parser')
    assert not body.select('details')
    assert body.select_one('strong').text == '標題'
    assert 'Hidden child content.' in body.text
    assert body.select_one('table td').text == '資料'
    assert '</details>' in body.select_one('pre').text
    assert body.select_one('a')['href'] == 'https://example.invalid/guide'
    document = Document(app.export_word(source))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '外層 標題' in text and '內層標題' in text and 'Hidden child content.' in text
    assert '</details>' in text and '<summary>literal</summary>' in text
    assert document.tables[0].cell(1, 1).text == '完整保留'
    pdf = Path(app.export_pdf(source))
    pdf_text = subprocess.check_output(['pdftotext', str(pdf), '-'], text=True)
    assert 'Hidden child content.' in pdf_text and '完整保留' in pdf_text
    assert '</details>' in pdf_text and '<summary>literal</summary>' in pdf_text


def test_commonmark_list_children_keep_exported_hierarchy(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = '''<details>
<summary>Nested plan</summary>

- Parent
  - [x] Child
    - Grandchild

    Child paragraph.

    ```text
    - literal code
    ```

1. Ordered parent
   1. Ordered child

> - Quoted parent
>   - Quoted child

</details>
'''
    body = BeautifulSoup(export_markdown.render_export_body(source, THEMES['Light']), 'html.parser')
    assert len(body.select('li > ul > li > ul > li')) == 1
    assert body.select_one('li > ol > li').get_text(strip=True) == 'Ordered child'
    assert body.select_one('blockquote li > ul > li').get_text(strip=True) == 'Quoted child'
    assert any(p.get_text(strip=True) == 'Child paragraph.' for p in body.select('li > p'))
    assert '- literal code' in body.select_one('li .codehilite').get_text()
    doc = Document(app.export_word(source))
    levels = [p._p.pPr.numPr.ilvl.val for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert {0, 1, 2} <= set(levels)
