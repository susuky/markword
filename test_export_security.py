"""Exercise offline export boundaries through real renderers, not mocked PDFs."""

import base64
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
import threading
import urllib.request

from docx import Document
from PIL import Image
import pytest
from weasyprint.urls import FatalURLFetchingError

import app
import mermaid_renderer
from mermaid_renderer import render_mermaid_png
from themes import THEMES


def data_url(mime, content):
    return f'data:{mime};base64,{base64.b64encode(content.encode()).decode()}'


SVG_REMOTE_IMAGE = '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><image href="https://example.invalid/nested.png" width="20" height="20"/></svg>'


@pytest.mark.parametrize('markdown', [
    '![image](https://example.invalid/image.png)',
    '<img src="http://127.0.0.1:9/private">',
    '<img src="http://[::1]:9/private">',
    '<img src="file:///tmp/markword-private.svg">',
    '<img src="ftp://example.invalid/image.png">',
    '<img src="HtTp&#58;//example.invalid/image.png">',
    '<base href="file:///tmp/"><img src="markword-private.svg">',
    '<base href="https://example.invalid/"><img src="//example.invalid/image.png">',
    '<style>@import url(https://example.invalid/private.css);</style>',
    '<p style="background-image:url(file:///tmp/markword-private.svg)">Body</p>',
    '<style>@font-face {font-family: probe; src: url(file:///tmp/private.ttf)} body {font-family: probe}</style>',
    SVG_REMOTE_IMAGE,
    f'<img src="{data_url("image/svg+xml", SVG_REMOTE_IMAGE)}">',
    # MIME labels do not make embedded SVG inert; WeasyPrint sniffs the bytes.
    f'<img src="{data_url("image/png", SVG_REMOTE_IMAGE)}">',
    f'<link rel="stylesheet" href="{data_url("text/css", "@import url(https://example.invalid/nested.css);")}">',
    '<link rel="attachment" href="file:///tmp/private.txt">',
    '<a rel="attachment" href="https://example.invalid/private.txt">Attachment</a>',
])
def test_pdf_denies_resources_before_opening_them(markdown, monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    opened = []
    original_open = urllib.request.OpenerDirector.open

    def guarded_open(self, request, *args, **kwargs):
        url = request.full_url if isinstance(request, urllib.request.Request) else request
        opened.append(url)
        if not url.startswith('data:'):
            raise OSError('Test prevented external resource access')
        return original_open(self, request, *args, **kwargs)

    monkeypatch.setattr(urllib.request.OpenerDirector, 'open', guarded_open)
    with pytest.raises(FatalURLFetchingError):
        app.export_pdf('# Resource boundary\n\n' + markdown)
    assert all(url.startswith('data:') for url in opened)
    assert list(tmp_path.iterdir()) == []  # failed exports leave no partial files


def test_pdf_keeps_embedded_images_styles_and_normal_links(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    write_pdf = app.HTML.write_pdf

    def uncompressed_pdf(self, *args, **kwargs):
        return write_pdf(self, *args, **kwargs, uncompressed_pdf=True)

    monkeypatch.setattr(app.HTML, 'write_pdf', uncompressed_pdf)
    image = BytesIO()
    Image.new('RGB', (20, 20), 'red').save(image, format='PNG')
    source = (
        '# 本機匯出\n\n[普通連結](https://example.invalid/)\n\n'
        f'<img src="data:image/png;base64,{base64.b64encode(image.getvalue()).decode()}">'
        '<style>h1 {color: green}</style>'
    )
    result = Path(app.export_pdf(source)).read_bytes()
    assert result.startswith(b'%PDF-')
    assert b'/Subtype /Image' in result
    assert b'https://example.invalid/' in result


@pytest.mark.parametrize('theme', ['Light', 'Dark'])
def test_mermaid_is_embedded_in_both_export_formats(theme, monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    # Disallow Python's old cloud-rendering path as well as PDF network fetches.
    def deny_urlopen(*args, **kwargs):
        pytest.fail('Export tried to call an external renderer')

    monkeypatch.setattr(urllib.request, 'urlopen', deny_urlopen)
    markdown = '# 公司架構\n\n```mermaid\nflowchart LR\n A[公司] --> B[客戶資料]\n```'
    pdf = Path(app.export_pdf(markdown, theme)).read_bytes()
    doc = Document(app.export_word(markdown, theme))
    assert pdf.startswith(b'%PDF-') and b'/Subtype /Image' in pdf
    assert len(doc.inline_shapes) == 1
    assert '公司架構' in doc.paragraphs[0].text
    assert all('could not be rendered' not in p.text for p in doc.paragraphs)


def test_mermaid_cannot_fetch_images_or_styles_from_directives():
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.end_headers()

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    origin = f'http://127.0.0.1:{server.server_port}'
    try:
        source = (
            '%%{init: {"securityLevel":"loose", "themeCSS":"@import url('
            + origin + '/private.css);"}}%%\nflowchart LR\n'
            + 'A@{ img: "' + origin + '/image.png", label: "Image", h: 60 } --> B[Done]'
        )
        # A diagram depending on a denied image cannot be drawn; retain its
        # source via the export fallback instead of fetching the missing image.
        assert render_mermaid_png(source, THEMES['Light']) is None
        assert requests == []
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


def test_invalid_mermaid_preserves_source_without_fetching(monkeypatch, tmp_path):
    monkeypatch.setattr(app, 'EXPORT_DIR', str(tmp_path))
    source = 'not-a-diagram <b>keep this source</b>'
    markdown = f'```mermaid\n{source}\n```'
    html = app._render_md_to_html_for_export(markdown, THEMES['Light'])
    assert 'could not be rendered' in html
    assert '&lt;b&gt;keep this source&lt;/b&gt;' in html
    doc = Document(app.export_word(markdown))
    assert source in '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 0


def test_large_mermaid_uses_source_fallback():
    source = 'flowchart LR\n' + ' --> '.join(f'N{i}[Node {i}]' for i in range(50))
    assert render_mermaid_png(source, THEMES['Light']) is None


def test_missing_mermaid_bundle_fails_without_remote_fallback(monkeypatch, tmp_path):
    monkeypatch.setattr(mermaid_renderer, 'MERMAID_BUNDLE', tmp_path / 'missing.js')
    with pytest.raises(RuntimeError, match='Local Mermaid bundle missing'):
        render_mermaid_png('flowchart LR\nA --> B', THEMES['Light'])


def test_legacy_preview_embeds_local_mermaid_bundle():
    from bs4 import BeautifulSoup

    preview = BeautifulSoup(app.render_preview('```mermaid\nflowchart LR\nA --> B\n```'), 'html.parser')
    document = BeautifulSoup(preview.iframe['srcdoc'], 'html.parser')
    assert document.find('script', src=True) is None
    assert any('mermaid.initialize' in script.text for script in document.find_all('script'))
